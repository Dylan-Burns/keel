"""The memo export of chapter 27.2: Markdown, DOCX, sources, and links."""

import importlib
import io
import json
import pathlib
import re
import types

import docx
import pytest

import export_samples as samples
from keel import audit
from keel import cache
from keel import cli
from keel import config
from keel import export

AAPL = samples.AAPL
AAPL_MDNA = samples.AAPL_MDNA
MSFT = samples.MSFT
UNCITED = samples.UNCITED
AAPL_URL = samples.AAPL_URL
MSFT_URL = samples.MSFT_URL
EVIL_URL = samples.EVIL_URL
QUESTION = samples.QUESTION
make_record = samples.make_record
audit_line = samples.audit_line


def test_export_markdown() -> None:
    memo = export.build_memo(make_record(), "req-1")
    text = export.render_markdown(memo)
    lines = text.splitlines()
    assert lines[0] == f"# {QUESTION}"
    assert "_Request req-1_  " in lines
    assert "_Model claude-sonnet-5-5, prompt 1.3.0, index 1.0.0_  " in lines
    # Numbers follow the first reference; the prose keeps its labels.
    assert (
        "Apple reported net sales of $416,161 million in FY2025 \\(period "
        "ended 2025-09-27\\)[^1]. Microsoft names supplier concentration"
        "[^2][^1]. See \\[note\\] for scope."
    ) in text
    assert "- Single-source suppliers in Asia.[^3]" in lines
    assert "- A claim with no valid citation. (unverified)" in lines
    assert "- Datacenter parts[^2]." in lines
    assert "**Confidence:** medium: 2 of 3 companies covered" in lines
    assert "| Company | Risk \\| theme |" in lines
    assert "| Apple Inc. | Suppliers \\<b\\>Asia\\</b\\>[^1] |" in lines
    assert "| Microsoft Corp |  |" in lines
    assert "- No 10-Q after Q3 FY2025." in lines
    assert "- AAPL: 10-K; FY2025" in lines
    assert "- Not covered: NVDA" in lines
    sources = lines[lines.index("## Sources") + 2 :]
    assert [line.split(":")[0] for line in sources if line] == [
        "[^1]",
        "[^2]",
        "[^3]",
        "[^4]",
    ]
    assert sources[0] == (
        "[^1]: Apple Inc. 10-K FY2025, Risk Factors, FY2025, period ended "
        '2025-09-27, filed 2025-10-31. Quote: "net sales of $416,161 '
        'million". Quote: "second quote". Chunk '
        f"{AAPL.replace('_', chr(92) + '_')}. SEC filing: <{AAPL_URL}>"
    )
    # A chunk from the retrieval list: its section title and link.
    assert sources[2].startswith(
        "[^3]: Apple Inc. 10-K FY2025, Management's Discussion and Analysis"
    )
    # A validated citation the text never references still gets a note.
    assert sources[3].startswith("[^4]: Microsoft Corp 10-K FY2025")
    assert "SEC filing" not in sources[3]


def test_export_docx(tmp_path: pathlib.Path) -> None:
    memo = export.build_memo(make_record(), "req-1")
    path = tmp_path / "memo.docx"
    path.write_bytes(export.render_docx(memo))
    document = docx.Document(str(path))
    assert document.core_properties.title == QUESTION
    texts = [p.text for p in document.paragraphs]
    assert texts[0] == QUESTION
    sources = texts[texts.index("Sources") + 1 :]
    assert [s.split(".")[0] for s in sources] == ["1", "2", "3", "4"]
    assert sources[0].startswith(
        "1. Apple Inc. 10-K FY2025, Risk Factors, FY2025, period ended "
        "2025-09-27"
    )
    assert sources[0].endswith(f"SEC filing: {AAPL_URL}")
    # The superscript markers in the direct answer follow citation order.
    answer = document.paragraphs[texts.index("Answer") + 1]
    marks = [r.text for r in answer.runs if r.font.superscript]
    assert marks == ["1", "2", ",1"]
    assert "(period ended 2025-09-27)" in answer.text
    # Every trusted source carries a real external link to www.sec.gov.
    links = [
        link.url
        for p in document.paragraphs[texts.index("Sources") + 1 :]
        for link in p.hyperlinks
    ]
    assert links == [AAPL_URL, MSFT_URL, AAPL_URL]
    targets = {
        rel.target_ref for rel in document.part.rels.values() if rel.is_external
    }
    assert targets == {AAPL_URL, MSFT_URL}
    table = document.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Company", "Risk | theme"]
    assert table.rows[1].cells[1].text == "Suppliers <b>Asia</b>1"
    bullets = [
        p.text for p in document.paragraphs if p.style.name == "List Bullet"
    ]
    assert "A claim with no valid citation. (unverified)" in bullets


def test_export_no_sec_link_for_untrusted_url(tmp_path: pathlib.Path) -> None:
    record = make_record(msft_url=EVIL_URL)
    memo = export.build_memo(record, "req-1")
    msft = next(s for s in memo.sources if s.chunk_id == MSFT)
    assert msft.sec_url is None
    text = export.render_markdown(memo)
    assert "evil.example" not in text
    path = tmp_path / "memo.docx"
    path.write_bytes(export.render_docx(memo))
    document = docx.Document(str(path))
    targets = [r.target_ref for r in document.part.rels.values()]
    assert not any("evil.example" in t for t in targets)
    assert all("evil.example" not in p.text for p in document.paragraphs)


@pytest.mark.parametrize(
    "url",
    [
        "http://www.sec.gov/x.htm",
        "https://www.sec.gov.evil.example/x.htm",
        "https://user@www.sec.gov/x.htm",
        "javascript:alert(1)",
    ],
)
def test_export_link_needs_the_sec_host(url: str) -> None:
    memo = export.build_memo(make_record(msft_url=url), "req-1")
    assert next(s for s in memo.sources if s.chunk_id == MSFT).sec_url is None


def test_export_escapes_model_text() -> None:
    record = make_record()
    record["question"] = "Risk\n<script>x</script>"
    record["answer"]["gaps"] = ["fake [^9] note"]
    text = export.render_markdown(export.build_memo(record, "req-1"))
    assert "<script>" not in text
    assert text.startswith("# Risk \\<script\\>x\\</script\\>\n")
    assert "- fake \\[^9\\] note" in text


def test_export_without_answer_or_table() -> None:
    record = make_record()
    record["answer"].update(
        comparison_table=None,
        company_findings=[],
        gaps=[],
        refusal={"reason": "outside the filings"},
    )
    record["citations_resolved"] = []
    record["retrieval"] = None
    record["question"] = ""
    memo = export.build_memo(record, "req-1")
    assert memo.title == "Answer req-1"
    text = export.render_markdown(memo)
    assert "## Comparison table" not in text
    assert f"- {export.NO_GAPS}" in text
    assert "**Declined:** outside the filings" in text
    # A marker with no known source still gets a numbered note.
    assert "[^1]: AAPL\\_10K\\_2025-09-27\\_\\_risk\\_factors" in text
    document = export.render_docx(memo)
    assert document[:2] == b"PK"
    record["answer"] = None
    with pytest.raises(export.ExportError, match="no model answer"):
        export.build_memo(record, "req-1")


def test_export_needs_the_docx_extra(monkeypatch: pytest.MonkeyPatch) -> None:
    real = importlib.import_module

    def missing(name: str) -> types.ModuleType:
        if name.startswith("docx"):
            raise ImportError(name)
        return real(name)

    monkeypatch.setattr(importlib, "import_module", missing)
    memo = export.build_memo(make_record(), "req-1")
    with pytest.raises(export.ExportError, match="export extra"):
        export.render(memo, export.DOCX)
    assert export.render(memo, export.MARKDOWN).startswith(b"# ")
    with pytest.raises(export.ExportError, match="unknown memo format"):
        export.render(memo, "pdf")


def store_settings(
    settings: config.Settings, tmp_path: pathlib.Path
) -> config.Settings:
    return settings.model_copy(
        update={
            "logs": config.LogsConfig(dir=tmp_path / "logs"),
            "cache": config.CacheConfig(answers_dir=tmp_path / "answers"),
        }
    )


def test_export_finds_the_answer_through_the_audit_record(
    settings: config.Settings, tmp_path: pathlib.Path
) -> None:
    local = store_settings(settings, tmp_path)
    answers = cache.AnswerCache(local.cache.answers_dir)
    key = cache.cache_key(QUESTION, "1.0.0", "1.3.0", "claude-sonnet-5-5")
    answers.put(key, make_record("first-1"))
    log = audit.AuditLog(local.logs.dir)
    assert log.write(audit_line("other-1"))
    # A later request served from the cache has its own request ID.
    assert log.write(audit_line("later-2"))
    found = export.find_record(local, "later-2")
    assert found["request_id"] == "first-1"
    memo = export.build_memo(found, "later-2")
    assert memo.details[:2] == (
        "Request later-2",
        "Answer first given in request first-1",
    )
    # A rotated file and the request that filled the entry.
    (local.logs.dir / audit.FILE_NAME).rename(
        local.logs.dir / "audit-2026-10-20.jsonl"
    )
    assert export.find_record(local, "later-2")["request_id"] == "first-1"
    assert export.find_record(local, "first-1")["request_id"] == "first-1"
    # A fallback answer never enters the cache.
    assert log.write(audit_line("fallback-3", model="claude-haiku-5-5"))
    with pytest.raises(export.ExportError, match="no stored answer"):
        export.find_record(local, "fallback-3")
    with pytest.raises(export.ExportError, match="not a request ID"):
        export.find_record(local, "../etc")


def test_export_never_borrows_another_answer(
    settings: config.Settings, tmp_path: pathlib.Path
) -> None:
    local = store_settings(settings, tmp_path)
    answers = cache.AnswerCache(local.cache.answers_dir)
    key = cache.cache_key(QUESTION, "1.0.0", "1.3.0", "claude-sonnet-5-5")
    # A later answer replaced the entry that request old-1 filled.
    answers.put(key, make_record("newer-9"))
    log = audit.AuditLog(local.logs.dir)
    assert log.write(audit_line("old-1", outcome="answered"))
    # A refused question with the same key never takes the cached answer.
    assert log.write(audit_line("refused-2", outcome="out_of_scope"))
    for request_id in ("old-1", "refused-2"):
        with pytest.raises(export.ExportError, match="no stored answer"):
            export.find_record(local, request_id)
    assert log.write(audit_line("newer-9", outcome="answered"))
    assert export.find_record(local, "newer-9")["request_id"] == "newer-9"
    # Without the scan, an ID with no audit line finds nothing.
    with pytest.raises(export.ExportError, match="no stored answer"):
        export.find_record(local, "newer-0", scan=False)


def test_export_drops_xml_invalid_characters() -> None:
    record = make_record()
    record["question"] = "Risk\x00 of\x0b supply"
    record["answer"]["gaps"] = ["gap\x1f text"]
    record["citations_resolved"][0]["filing_name"] = "Microsoft\x08 Corp"
    memo = export.build_memo(record, "req-1")
    assert memo.title == "Risk of supply"
    document = docx.Document(io.BytesIO(export.render_docx(memo)))
    assert "gap text" in [p.text for p in document.paragraphs]


def test_export_from_a_saved_result(
    settings: config.Settings, tmp_path: pathlib.Path
) -> None:
    saved = tmp_path / "result.json"
    saved.write_text(json.dumps(make_record("saved-1")))
    assert export.find_record(settings, "saved-1", saved)["question"]
    with pytest.raises(export.ExportError, match="holds request saved-1"):
        export.find_record(settings, "other-1", saved)
    saved.write_text("[1]")
    with pytest.raises(export.ExportError, match="not hold an answer"):
        export.find_record(settings, "saved-1", saved)
    with pytest.raises(export.ExportError, match="cannot read"):
        export.find_record(settings, "saved-1", tmp_path / "absent.json")


def test_export_command(
    settings: config.Settings,
    tmp_path: pathlib.Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    local = store_settings(settings, tmp_path)
    monkeypatch.setattr(config, "load_settings", lambda _path=None: local)
    saved = tmp_path / "result.json"
    saved.write_text(json.dumps(make_record("cli-1")))
    out = tmp_path / "memos" / "memo.docx"
    argv = ["export", "cli-1", "--record", str(saved)]
    assert cli.main([*argv, "--format", "docx", "--out", str(out)]) == 0
    assert "4 sources" in capsys.readouterr().out
    assert len(docx.Document(str(out)).tables) == 1
    monkeypatch.chdir(tmp_path)
    assert cli.main(argv) == 0
    assert (tmp_path / "memo-cli-1.md").read_text().startswith("# How do")
    # No stored answer: a KeelError exit.
    assert cli.main(["export", "absent-1"]) == 2
    blocked = tmp_path / "file"
    blocked.write_text("")
    assert cli.main([*argv, "--out", str(blocked / "memo.md")]) == 2


INJECTION = (
    "[click](https://evil.example) ![x](https://evil.example/a.png) "
    "<https://evil.example> # heading\n- list *bold* `code` www.evil.example "
    "a@evil.example | cell"
)
# A metacharacter without a backslash before it.
LIVE = re.compile(r"(?<!\\)[\[\]()<>!*`|@]")


def test_export_markdown_injection_is_inert() -> None:
    record = make_record()
    answer = record["answer"]
    answer["gaps"] = [INJECTION]
    answer["company_findings"][0]["findings"][0]["point"] = INJECTION
    answer["comparison_table"]["rows"][0][1] = INJECTION
    record["citations_resolved"][0]["quote"] = INJECTION
    text = export.render_markdown(export.build_memo(record, "req-1"))
    assert "https://evil" not in text
    assert "www.evil" not in text
    lines = [line for line in text.splitlines() if "evil" in line]
    assert len(lines) == 4
    for line in lines:
        # The memo's own syntax: the list or note prefix and footnote refs.
        body = re.sub(r"\[\^\d+\](: )?", "", line).removeprefix("- ")
        body = re.sub(r" SEC filing: <https://www\.sec\.gov/\S+>$", "", body)
        body = body.strip("| ") if line.startswith("|") else body
        cells = body.split(" | ") if line.startswith("|") else [body]
        for cell in cells:
            assert not LIVE.search(cell), cell
    # The newline of the claim does not start a block of its own.
    assert not any(line.startswith(("# heading", "- list")) for line in lines)
    gap = next(line for line in lines if line.startswith("- \\[click"))
    assert "\\# heading \\- list" not in gap  # a hyphen mid-text stays
    assert "\\# heading - list \\*bold\\*" in gap


def test_export_sec_link_is_encoded() -> None:
    url = "https://www.sec.gov/Archives/a>b)c[d].htm"
    memo = export.build_memo(make_record(msft_url=url), "req-1")
    msft = next(s for s in memo.sources if s.chunk_id == MSFT)
    assert msft.sec_url == ("https://www.sec.gov/Archives/a%3Eb%29c%5Bd%5D.htm")
    text = export.render_markdown(memo)
    assert f"<{msft.sec_url}>" in text
    assert "a>b" not in text
    document = docx.Document(io.BytesIO(export.render_docx(memo)))
    links = [link.url for p in document.paragraphs for link in p.hyperlinks]
    assert msft.sec_url in links
    assert all(">" not in link and ")" not in link for link in links)
