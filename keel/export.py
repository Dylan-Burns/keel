"""Memo export (chapter 27.2): 1 stored answer as a Markdown or DOCX memo.

`keel export REQUEST_ID` finds the stored answer of a request ID and writes
a memo: the question as the title, the direct answer, the findings per
company, the comparison table, the gaps, the coverage, and numbered
footnotes with the filing, the section, the fiscal label, the period end,
and the SEC link.

The lookup reads, in order: a saved result JSON (`--record`, the footer
Export JSON or `keel ask --json`), the audit record of the request ID (its
question, index version, prompt version, and model give the answer cache
key of decision D-29), and the answer cache entry that the request filled.

The text of the answer stays as the model wrote it and the validator kept
it, so every number keeps its fiscal label and period end date. Each
inline marker `[chunk_id]` becomes a footnote number; numbers follow the
order of first reference in the memo, and a validated citation that the
text never references comes after them. A link renders only for a URL that
passes the `www.sec.gov` check of control S-05 (E-D-28), checked again here
because a saved result JSON comes from outside the pipeline.

The DOCX writer needs python-docx (the `export` extra). It is imported on
first use, so the core install runs without it. python-docx has no
footnote API, so the DOCX memo carries superscript numbers in the text
and the numbered notes in a `Sources` section at the end (endnotes).

Each written memo appends 1 line to `<logs.dir>/exports.jsonl` (the time,
the request ID, the format, and the user), which the pilot metrics count
(chapter 27.6).
"""

from __future__ import annotations

import argparse
from collections.abc import Iterable
from collections.abc import Mapping
import dataclasses
import datetime
import importlib
import io
import json
import logging
import os
import pathlib
import re
import sys
from typing import TYPE_CHECKING
from typing import Any
import urllib.parse

from keel import audit
from keel import cache
from keel import citations
from keel import config
from keel import errors
from keel.parsers import form_10k_10q

if TYPE_CHECKING:
    from docx import document as docx_document
    from docx.text import paragraph as docx_paragraph

logger = logging.getLogger(__name__)

MARKDOWN = "markdown"
DOCX = "docx"
FORMATS = (MARKDOWN, DOCX)
SUFFIXES = {MARKDOWN: ".md", DOCX: ".docx"}
MEDIA_TYPES = {
    MARKDOWN: "text/markdown; charset=utf-8",
    DOCX: (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    ),
}
REQUEST_ID = re.compile(r"^[A-Za-z0-9._:-]{1,128}$")
NO_GAPS = "The answer names no gaps."
EXTRA_HINT = (
    "the DOCX memo needs python-docx: install the export extra "
    "(uv sync --extra export)"
)
LINK_COLOR = "0563C1"
# Every Markdown metacharacter of untrusted text gets a backslash; a hyphen
# or a plus sign only at the start of a field, where it would start a list.
_MARKDOWN_SPECIAL = re.compile(r"([\\`*_\[\]()#!|<>{}&~@])")
_SCHEME = re.compile(r"\b([A-Za-z][A-Za-z0-9+.\-]{0,31}):(?=\S)")
_WWW = re.compile(r"\b(www)\.", re.IGNORECASE)
_BLOCK_START = re.compile(r"^(\d+)([.)])|^([-+=])")
# The characters left unencoded in a link URL; every other one, such as
# `<>[]()` and whitespace, is percent-encoded.
URL_SAFE = "/:?=&%#.-_~"
_URL_FORBIDDEN = re.compile(r"[<>\[\]()`\s\\]")
# The code point ranges that XML 1.0 (and so a DOCX file) can hold.
_XML_VALID = (
    (0x20, 0xD7FF),
    (0xE000, 0xFFFD),
    (0x10000, 0x10FFFF),
)


def _xml_char(char: str) -> bool:
    """Tell whether XML 1.0 can hold 1 character."""
    point = ord(char)
    return char in "\t\n\r" or any(lo <= point <= hi for lo, hi in _XML_VALID)


_MARKER = re.compile(r"\s*\[([^\[\]\n]{1,400})\]")

# A run of memo text: strings and footnote numbers in reading order.
Text = tuple[str | int, ...]


class ExportError(errors.KeelError):
    """No stored answer to export, or the memo cannot be written."""


# The export log of the pilot metrics (chapter 27.6): 1 line per memo.
LOG_NAME = "exports.jsonl"
CLI_USER = "cli"


def log_export(
    directory: pathlib.Path, request_id: str, fmt: str, user_name: str
) -> bool:
    """Append 1 line to `<logs.dir>/exports.jsonl` for a written memo.

    The line holds the time, the request ID, the format, and the user; no
    question, no answer text, and no token. The file is readable by the
    owner only, like the audit log. A write failure is logged and never
    fails the export.

    Args:
        directory: The logs directory.
        request_id: The request ID of the exported answer.
        fmt: `markdown` or `docx`.
        user_name: The user of the API route, or `cli` for `keel export`.

    Returns:
        True when the line was written; False after a logged failure.
    """
    line = {
        "time": datetime.datetime.now(tz=datetime.UTC).isoformat(),
        "request_id": request_id,
        "format": fmt,
        "user_name": user_name,
    }
    try:
        text = json.dumps(line, ensure_ascii=False)
        directory.mkdir(parents=True, exist_ok=True)
        handle = os.open(
            directory / LOG_NAME,
            os.O_WRONLY | os.O_APPEND | os.O_CREAT,
            audit.FILE_MODE,
        )
        with os.fdopen(handle, "a", encoding="utf-8") as stream:
            # Also tighten an older file; the stream closes on a failure.
            os.fchmod(stream.fileno(), audit.FILE_MODE)
            stream.write(text + "\n")
    except (OSError, TypeError, ValueError) as error:
        logger.error("export log write failed: %s", type(error).__name__)
        return False
    return True


@dataclasses.dataclass(frozen=True)
class Source:
    """1 numbered footnote: 1 cited chunk.

    Attributes:
        number: The footnote number.
        chunk_id: The chunk ID.
        filing_name: Such as `Apple Inc. 10-K FY2025`.
        section_title: Such as `Risk Factors`.
        fiscal_label: Such as `FY2025`.
        period_end: The period end date.
        filing_date: The filing date, or an empty string.
        sec_url: The trusted SEC URL, or None.
        quotes: The quotes the citations of the chunk carry.
    """

    number: int
    chunk_id: str
    filing_name: str
    section_title: str
    fiscal_label: str
    period_end: str
    filing_date: str
    sec_url: str | None
    quotes: tuple[str, ...]


@dataclasses.dataclass(frozen=True)
class Finding:
    """1 finding line.

    Attributes:
        text: The point with its footnote numbers.
        unverified: True when no citation was valid.
    """

    text: Text
    unverified: bool


@dataclasses.dataclass(frozen=True)
class Company:
    """The findings of 1 company.

    Attributes:
        heading: Such as `Apple Inc. (AAPL)`.
        findings: The findings.
    """

    heading: str
    findings: tuple[Finding, ...]


@dataclasses.dataclass(frozen=True)
class Memo:  # pylint: disable=too-many-instance-attributes
    """The parts of a memo, shared by the Markdown and DOCX writers.

    Attributes:
        title: The question.
        request_id: The request ID.
        details: The request, model, version, and data lines.
        direct_answer: The direct answer.
        confidence: Such as `medium: 2 of 3 companies covered`.
        refusal: The refusal reason, or None.
        companies: The findings per company.
        columns: The comparison table headers (empty without a table).
        rows: The comparison table rows.
        gaps: The gaps.
        coverage: The coverage lines.
        sources: The footnotes in number order.
    """

    title: str
    request_id: str
    details: tuple[str, ...]
    direct_answer: Text
    confidence: str
    refusal: str | None
    companies: tuple[Company, ...]
    columns: tuple[str, ...]
    rows: tuple[tuple[Text, ...], ...]
    gaps: tuple[str, ...]
    coverage: tuple[str, ...]
    sources: tuple[Source, ...]


def _trusted(url: object) -> str | None:
    """A URL fit for a link, or None (control S-05, E-D-28).

    The URL must pass the `www.sec.gov` check of the parser, then it is
    percent-encoded so no character can end a Markdown `<...>` link, and
    the encoded URL must pass the check again.

    Args:
        url: The stored URL.

    Returns:
        The encoded trusted URL, or None.
    """
    if not isinstance(url, str) or not form_10k_10q.url_trusted(url):
        return None
    encoded = urllib.parse.quote(url, safe=URL_SAFE)
    if _URL_FORBIDDEN.search(encoded) or not form_10k_10q.url_trusted(encoded):
        return None
    return encoded


def _flat(text: object) -> str:
    """Collapse the whitespace of a field, so no newline starts a block.

    Characters that XML cannot hold are dropped, so the DOCX writer never
    fails on them.
    """
    return " ".join("".join(c for c in str(text) if _xml_char(c)).split())


def _source_facts(record: Mapping[str, Any]) -> dict[str, dict[str, Any]]:
    """The facts of every chunk the record names, by chunk ID."""
    facts: dict[str, dict[str, Any]] = {}
    retrieval = record.get("retrieval") or {}
    for chunk in retrieval.get("chunks") or []:
        section = str(chunk.get("section", ""))
        facts[str(chunk.get("chunk_id"))] = {
            "filing_name": _flat(chunk.get("filing_name", "")),
            "section_title": citations.SECTION_TITLES.get(section, section),
            "fiscal_label": _flat(chunk.get("fiscal_label", "")),
            "period_end": _flat(chunk.get("period_end", "")),
            "filing_date": "",
            "sec_url": _trusted(chunk.get("sec_url")),
            "quotes": [],
        }
    for cited in record.get("citations_resolved") or []:
        chunk_id = str(cited.get("chunk_id"))
        known = facts.get(chunk_id, {})
        quotes = list(known.get("quotes", []))
        quote = _flat(cited.get("quote", ""))
        if quote and quote not in quotes:
            quotes.append(quote)
        facts[chunk_id] = {
            "filing_name": _flat(cited.get("filing_name", "")),
            "section_title": _flat(cited.get("section_title", "")),
            "fiscal_label": _flat(cited.get("fiscal_label", "")),
            "period_end": _flat(cited.get("period_end", "")),
            "filing_date": _flat(cited.get("filing_date", "")),
            "sec_url": _trusted(cited.get("sec_url")),
            "quotes": quotes,
        }
    return facts


class _Numbering:
    """Give each cited chunk its footnote number at first reference."""

    def __init__(self) -> None:
        self.numbers: dict[str, int] = {}

    def number(self, chunk_id: str) -> int:
        """Return the number of a chunk, giving the next one when new."""
        if chunk_id not in self.numbers:
            self.numbers[chunk_id] = len(self.numbers) + 1
        return self.numbers[chunk_id]

    def text(self, raw: str, extra: Iterable[str] = ()) -> Text:
        """Replace the inline markers of a text with footnote numbers.

        Args:
            raw: The text with `[chunk_id, ...]` markers.
            extra: Chunk IDs to reference at the end when the text does
                not already reference them (the citations of a finding).

        Returns:
            The text pieces and the numbers.
        """
        raw = _flat(raw)
        pieces: list[str | int] = []
        seen: list[str] = []
        position = 0
        for match in _MARKER.finditer(raw):
            ids = [p.strip() for p in re.split(r"[,;]", match.group(1))]
            if not all(citations.CHUNK_ID.match(i) for i in ids):
                continue  # brackets that are not a citation marker
            pieces.append(raw[position : match.start()])
            for chunk_id in ids:
                pieces.append(self.number(chunk_id))
                seen.append(chunk_id)
            position = match.end()
        pieces.append(raw[position:])
        for chunk_id in extra:
            if chunk_id not in seen:
                pieces.append(self.number(chunk_id))
                seen.append(chunk_id)
        return tuple(p for p in pieces if p != "")


def _details(record: Mapping[str, Any], request_id: str) -> tuple[str, ...]:
    meta = record.get("meta") or {}
    answer = record.get("answer") or {}
    lines = [f"Request {request_id}"]
    first = record.get("request_id")
    if first and first != request_id:
        lines.append(f"Answer first given in request {first}")
    parts = [
        f"{label} {meta[key]}"
        for label, key in (
            ("Model", "model"),
            ("prompt", "prompt_version"),
            ("index", "index_version"),
        )
        if meta.get(key)
    ]
    if parts:
        lines.append(", ".join(parts))
    data_as_of = (answer.get("coverage") or {}).get("data_as_of")
    if data_as_of:
        lines.append(f"Data as of: {data_as_of}")
    return tuple(_flat(line) for line in lines)


def _coverage(answer: Mapping[str, Any]) -> tuple[str, ...]:
    coverage = answer.get("coverage") or {}
    lines = []
    for used in coverage.get("companies_used") or []:
        forms = ", ".join(used.get("forms") or [])
        periods = ", ".join(used.get("fiscal_periods") or [])
        ticker = used.get("ticker", "")
        lines.append(f"{ticker}: {forms}; {periods}")
    missing = coverage.get("companies_missing") or []
    if missing:
        lines.append("Not covered: " + ", ".join(missing))
    return tuple(_flat(line) for line in lines)


def build_memo(record: Mapping[str, Any], request_id: str) -> Memo:
    """Build the memo parts of a stored answer record.

    Args:
        record: The answer record of chapter 14.2 (the `final` event).
        request_id: The request ID the memo names.

    Returns:
        The memo.

    Raises:
        ExportError: The record holds no model answer.
    """
    answer = record.get("answer")
    if not isinstance(answer, Mapping):
        outcome = record.get("outcome")
        raise ExportError(
            f"request {request_id} has no model answer to export "
            f"(outcome {outcome})"
        )
    numbering = _Numbering()
    direct = numbering.text(str(answer.get("direct_answer", "")))
    companies = []
    for company in answer.get("company_findings") or []:
        findings = tuple(
            Finding(
                text=numbering.text(
                    str(f.get("point", "")), f.get("citations") or []
                ),
                unverified=bool(f.get("unverified")),
            )
            for f in company.get("findings") or []
        )
        name, ticker = company.get("company", ""), company.get("ticker", "")
        heading = _flat(f"{name} ({ticker})")
        companies.append(Company(heading=heading, findings=findings))
    table = answer.get("comparison_table") or {}
    columns = tuple(_flat(c) for c in table.get("columns") or [])
    rows = tuple(
        tuple(numbering.text(str(cell)) for cell in row)
        for row in table.get("rows") or []
    )
    facts = _source_facts(record)
    for cited in record.get("citations_resolved") or []:
        numbering.number(str(cited.get("chunk_id")))
    sources = []
    for chunk_id, number in numbering.numbers.items():
        fact = facts.get(chunk_id, {})
        sources.append(
            Source(
                number=number,
                chunk_id=chunk_id,
                filing_name=fact.get("filing_name") or chunk_id,
                section_title=fact.get("section_title", ""),
                fiscal_label=fact.get("fiscal_label", ""),
                period_end=fact.get("period_end", ""),
                filing_date=fact.get("filing_date", ""),
                sec_url=fact.get("sec_url"),
                quotes=tuple(fact.get("quotes", [])),
            )
        )
    refusal = answer.get("refusal") or {}
    confidence = str(answer.get("confidence", ""))
    reason = _flat(answer.get("confidence_reason", ""))
    question = _flat(record.get("question") or "")
    return Memo(
        title=question or f"Answer {request_id}",
        request_id=request_id,
        details=_details(record, request_id),
        direct_answer=direct,
        confidence=f"{confidence}: {reason}" if reason else confidence,
        refusal=_flat(refusal.get("reason", "")) if refusal else None,
        companies=tuple(companies),
        columns=columns,
        rows=rows,
        gaps=tuple(_flat(g) for g in answer.get("gaps") or []),
        coverage=_coverage(answer),
        sources=tuple(sources),
    )


def source_note(source: Source) -> str:
    """The text of 1 footnote without the link.

    Args:
        source: The source.

    Returns:
        Such as `Apple Inc. 10-K FY2025, Risk Factors, FY2025, period
        ended 2025-09-27, filed 2025-10-31. Quote: "..."`.
    """
    parts = [source.filing_name]
    if source.section_title:
        parts.append(source.section_title)
    if source.fiscal_label:
        parts.append(source.fiscal_label)
    if source.period_end:
        parts.append(f"period ended {source.period_end}")
    if source.filing_date:
        parts.append(f"filed {source.filing_date}")
    note = ", ".join(parts) + "."
    for quote in source.quotes:
        note += f' Quote: "{quote}".'
    if source.filing_name != source.chunk_id:
        note += f" Chunk {source.chunk_id}."
    return note


def _escape(text: str, start: bool = True) -> str:
    """Make untrusted text inert Markdown inline text (control S-05).

    Newlines collapse, every metacharacter that can open a link, an image,
    an autolink, HTML, emphasis, code, or a table cell gets a backslash,
    the colon after a scheme and the dot after `www` are escaped so no
    bare URL becomes an autolink, and a list or heading marker at the
    start of the field is escaped.

    Args:
        text: The untrusted text.
        start: True when the text starts a line of the memo.

    Returns:
        The escaped text.
    """
    flat = _flat(text)
    escaped = _MARKDOWN_SPECIAL.sub(r"\\\1", flat)
    escaped = _SCHEME.sub(r"\1\\:", escaped)
    escaped = _WWW.sub(r"\1\\.", escaped)
    if start:
        escaped = _BLOCK_START.sub(_block_escape, escaped, count=1)
    return escaped


def _block_escape(match: re.Match[str]) -> str:
    if match.group(3):
        return "\\" + match.group(3)
    return match.group(1) + "\\" + match.group(2)


def _cell(text: str) -> str:
    return _escape(text)


def _markdown_text(text: Text, start: bool = True) -> str:
    out = []
    for k, piece in enumerate(text):
        if isinstance(piece, int):
            out.append(f"[^{piece}]")
        else:
            out.append(_escape(piece, start=start and k == 0))
    return "".join(out)


def render_markdown(memo: Memo) -> str:
    """Write the memo as Markdown with numbered footnotes.

    Args:
        memo: The memo.

    Returns:
        The Markdown text.
    """
    lines = [f"# {_escape(memo.title)}", ""]
    lines += [f"_{_escape(line)}_  " for line in memo.details]
    lines += ["", "## Answer", "", _markdown_text(memo.direct_answer), ""]
    if memo.refusal:
        lines += [f"**Declined:** {_escape(memo.refusal)}", ""]
    lines += [f"**Confidence:** {_escape(memo.confidence)}", ""]
    if memo.companies:
        lines += ["## Findings by company", ""]
        for company in memo.companies:
            lines += [f"### {_escape(company.heading)}", ""]
            for finding in company.findings:
                mark = " (unverified)" if finding.unverified else ""
                lines.append(f"- {_markdown_text(finding.text)}{mark}")
            lines.append("")
    if memo.columns:
        lines += ["## Comparison table", ""]
        lines.append("| " + " | ".join(_cell(c) for c in memo.columns) + " |")
        lines.append("| " + " | ".join("---" for _ in memo.columns) + " |")
        for row in memo.rows:
            cells = [_cell_text(cell) for cell in row]
            cells = (cells + [""] * len(memo.columns))[: len(memo.columns)]
            lines.append("| " + " | ".join(cells) + " |")
        lines.append("")
    lines += ["## Gaps", ""]
    lines += [f"- {_escape(g)}" for g in memo.gaps] or [f"- {NO_GAPS}"]
    lines.append("")
    if memo.coverage:
        lines += ["## Coverage", ""]
        lines += [f"- {_escape(line)}" for line in memo.coverage]
        lines.append("")
    if memo.sources:
        lines += ["## Sources", ""]
        for source in memo.sources:
            note = _escape(source_note(source))
            if source.sec_url:
                note += f" SEC filing: <{source.sec_url}>"
            lines.append(f"[^{source.number}]: {note}")
        lines.append("")
    return "\n".join(lines)


def _cell_text(text: Text) -> str:
    return _markdown_text(text, start=False)


def _docx_modules() -> dict[str, Any]:
    """Import python-docx on first use (the `export` extra)."""
    try:
        return {
            "docx": importlib.import_module("docx"),
            "oxml": importlib.import_module("docx.oxml"),
            "ns": importlib.import_module("docx.oxml.ns"),
            "constants": importlib.import_module("docx.opc.constants"),
        }
    except ImportError as error:
        raise ExportError(EXTRA_HINT) from error


def _add_runs(paragraph: docx_paragraph.Paragraph, text: Text) -> None:
    """Add the text, with each footnote number as a superscript run."""
    previous_number = False
    for piece in text:
        if isinstance(piece, int):
            label = f",{piece}" if previous_number else str(piece)
            paragraph.add_run(label).font.superscript = True
            previous_number = True
        else:
            paragraph.add_run(piece)
            previous_number = False


def _add_hyperlink(
    paragraph: docx_paragraph.Paragraph,
    url: str,
    label: str,
    modules: Mapping[str, Any],
) -> None:
    """Add an external link run (python-docx has no hyperlink writer)."""
    oxml, qn = modules["oxml"], modules["ns"].qn
    relation = paragraph.part.relate_to(
        url,
        modules["constants"].RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    link = oxml.OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    link_run = oxml.OxmlElement("w:r")
    properties = oxml.OxmlElement("w:rPr")
    color = oxml.OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    underline = oxml.OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    link_run.append(properties)
    text = oxml.OxmlElement("w:t")
    text.text = label
    link_run.append(text)
    link.append(link_run)
    # pylint: disable-next=protected-access
    paragraph._p.append(link)


def _docx_findings(document: docx_document.Document, memo: Memo) -> None:
    """Add the findings per company as bullet lists."""
    if not memo.companies:
        return
    document.add_heading("Findings by company", level=1)
    for company in memo.companies:
        document.add_heading(company.heading, level=2)
        for finding in company.findings:
            item = document.add_paragraph(style="List Bullet")
            _add_runs(item, finding.text)
            if finding.unverified:
                item.add_run(" (unverified)").italic = True


def _docx_table(document: docx_document.Document, memo: Memo) -> None:
    """Add the comparison table with a bold header row."""
    if not memo.columns:
        return
    document.add_heading("Comparison table", level=1)
    table = document.add_table(rows=1 + len(memo.rows), cols=len(memo.columns))
    table.style = "Table Grid"
    for k, column in enumerate(memo.columns):
        table.cell(0, k).paragraphs[0].add_run(column).bold = True
    for r, row in enumerate(memo.rows, start=1):
        for k, cell in enumerate(row[: len(memo.columns)]):
            _add_runs(table.cell(r, k).paragraphs[0], cell)


def _docx_sources(
    document: docx_document.Document,
    memo: Memo,
    modules: Mapping[str, Any],
) -> None:
    """Add the numbered endnotes, each with its SEC link when trusted."""
    if not memo.sources:
        return
    document.add_heading("Sources", level=1)
    for source in memo.sources:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{source.number}. ").bold = True
        paragraph.add_run(source_note(source))
        if source.sec_url:
            paragraph.add_run(" SEC filing: ")
            _add_hyperlink(paragraph, source.sec_url, source.sec_url, modules)


def render_docx(memo: Memo) -> bytes:
    """Write the memo as a DOCX file with numbered endnotes and SEC links.

    Args:
        memo: The memo.

    Returns:
        The DOCX bytes.

    Raises:
        ExportError: python-docx is not installed.
    """
    modules = _docx_modules()
    document: docx_document.Document = modules["docx"].Document()
    document.core_properties.title = memo.title[:255]
    document.add_heading(memo.title, level=0)
    for line in memo.details:
        document.add_paragraph().add_run(line).italic = True
    document.add_heading("Answer", level=1)
    _add_runs(document.add_paragraph(), memo.direct_answer)
    if memo.refusal:
        paragraph = document.add_paragraph()
        paragraph.add_run("Declined: ").bold = True
        paragraph.add_run(memo.refusal)
    paragraph = document.add_paragraph()
    paragraph.add_run("Confidence: ").bold = True
    paragraph.add_run(memo.confidence)
    _docx_findings(document, memo)
    _docx_table(document, memo)
    document.add_heading("Gaps", level=1)
    for gap in memo.gaps or (NO_GAPS,):
        document.add_paragraph(gap, style="List Bullet")
    if memo.coverage:
        document.add_heading("Coverage", level=1)
        for line in memo.coverage:
            document.add_paragraph(line, style="List Bullet")
    _docx_sources(document, memo, modules)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _audit_entry(
    directory: pathlib.Path, request_id: str
) -> dict[str, Any] | None:
    """The audit record of a request ID, newest file first."""
    files = [directory / audit.FILE_NAME]
    files += sorted(directory.glob("audit-*.jsonl"), reverse=True)
    needle = json.dumps(request_id)
    for path in files:
        if not path.is_file():
            continue
        found = None
        with path.open(encoding="utf-8") as stream:
            for line in stream:
                if needle not in line:
                    continue
                try:
                    entry = json.loads(line)
                except ValueError:
                    continue
                if isinstance(entry, dict) and (
                    entry.get("request_id") == request_id
                ):
                    found = entry
        if found is not None:
            return found
    return None


def _from_cache(
    settings: config.Settings, request_id: str, scan: bool
) -> dict[str, Any] | None:
    """The answer cache entry of a request ID, through its audit record.

    The entry under the key of the audit record belongs to the request
    only when the request filled it or was served from it (outcome
    `cached`); a refused or failed request, or an entry a later answer
    replaced, never borrows another answer.
    """
    answers = cache.AnswerCache(settings.cache.answers_dir)
    entry = _audit_entry(settings.logs.dir, request_id)
    if entry is not None:
        parts = [
            entry.get(k)
            for k in ("question", "index_version", "prompt_version", "model")
        ]
        if all(isinstance(p, str) for p in parts):
            key = cache.cache_key(*[str(p) for p in parts])
            record = answers.get(key)
            if record is not None and (
                record.get("request_id") == request_id
                or entry.get("outcome") == "cached"
            ):
                return record
        logger.info("no answer cache entry for the audit record")
    return answers.find(request_id) if scan else None


def find_record(
    settings: config.Settings,
    request_id: str,
    record_path: pathlib.Path | None = None,
    scan: bool = True,
) -> dict[str, Any]:
    """Find the stored answer of a request ID.

    Args:
        settings: The loaded settings (`logs.dir`, `cache.answers_dir`).
        request_id: The request ID.
        record_path: A saved result JSON, or None to read the audit log
            and the answer cache.
        scan: Read every cache entry for the request ID when the audit
            record gives none; the API turns it off to bound the work.

    Returns:
        The answer record.

    Raises:
        ExportError: The request ID is malformed, the file does not hold
            the request, or no stored answer exists.
    """
    if not REQUEST_ID.match(request_id):
        raise ExportError(f"not a request ID: {request_id!r}")
    if record_path is not None:
        try:
            record = json.loads(record_path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as error:
            raise ExportError(f"cannot read {record_path}: {error}") from None
        if not isinstance(record, dict):
            raise ExportError(f"{record_path} does not hold an answer record")
        if record.get("request_id") != request_id:
            held = record.get("request_id")
            raise ExportError(
                f"{record_path} holds request {held}, not {request_id}"
            )
        return record
    found = _from_cache(settings, request_id, scan)
    if found is None:
        raise ExportError(
            f"no stored answer for request {request_id}: the audit log and "
            "the answer cache hold none (only a primary model answer of a "
            "real provider enters the cache)"
        )
    return found


def render(memo: Memo, fmt: str) -> bytes:
    """Write a memo in a format.

    Args:
        memo: The memo.
        fmt: `markdown` or `docx`.

    Returns:
        The file bytes.

    Raises:
        ExportError: An unknown format, or python-docx is missing.
    """
    if fmt == MARKDOWN:
        return render_markdown(memo).encode("utf-8")
    if fmt == DOCX:
        return render_docx(memo)
    raise ExportError(f"unknown memo format {fmt!r}")


def file_name(request_id: str, fmt: str) -> str:
    """The default file name of a memo.

    Args:
        request_id: The request ID (checked by `find_record`).
        fmt: `markdown` or `docx`.

    Returns:
        Such as `memo-abc123.docx`.
    """
    return f"memo-{request_id}{SUFFIXES[fmt]}"


def add_arguments(parser: argparse.ArgumentParser) -> None:
    """Add the arguments of `keel export` to its parser.

    Args:
        parser: The `export` subcommand parser.
    """
    parser.add_argument("request_id", help="the request ID of the answer")
    parser.add_argument(
        "--format",
        dest="memo_format",
        choices=FORMATS,
        default=MARKDOWN,
        help="markdown (default) or docx (the export extra)",
    )
    parser.add_argument(
        "--out",
        type=pathlib.Path,
        default=None,
        help="the memo file (default: memo-<request_id>.md or .docx)",
    )
    parser.add_argument(
        "--record",
        type=pathlib.Path,
        default=None,
        help="a saved result JSON instead of the audit log and the cache",
    )


def run(settings: config.Settings, args: argparse.Namespace) -> int:
    """Run `keel export`: find the answer, write the memo, print its path.

    Args:
        settings: The loaded settings.
        args: The parsed arguments of the `export` command.

    Returns:
        0 after the memo is written.

    Raises:
        ExportError: No stored answer, python-docx is missing, or the file
            cannot be written.
    """
    record = find_record(settings, args.request_id, args.record)
    memo = build_memo(record, args.request_id)
    data = render(memo, args.memo_format)
    out: pathlib.Path = args.out or pathlib.Path(
        file_name(args.request_id, args.memo_format)
    )
    try:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(data)
    except OSError as error:
        raise ExportError(f"cannot write {out}: {error}") from None
    log_export(settings.logs.dir, args.request_id, args.memo_format, CLI_USER)
    sys.stdout.write(
        f"memo {out} ({len(memo.sources)} sources, {len(data)} bytes)\n"
    )
    return 0
